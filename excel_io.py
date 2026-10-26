"""Reads the input workbook (NIT, VALOR, LIQUIDACION) and writes results."""
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

HEADER_ALIASES = {
    "nit": ("nit", "proveedor", "id poseedor"),
    "amount": ("valor", "total", "monto"),
    "liquidation": ("liquidacion", "liquidación"),
}


@dataclass
class PoRow:
    row: int
    nit: str
    amount: float
    liquidation: str
    po_number: str = ""
    status: str = "PENDIENTE"
    message: str = ""
    net_value: float | None = None  # value read from SAP before saving


def _match(header: str) -> str | None:
    h = (header or "").strip().lower()
    for key, aliases in HEADER_ALIASES.items():
        if any(a in h for a in aliases):
            return key
    return None


def _clean_id(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().replace(".", "").replace(",", "") if isinstance(value, (int, float)) else str(value).strip()


def _to_amount(value) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value or "").strip().replace("$", "").replace(" ", "")
    # Colombian format: 11.563.998 or 11.563.998,50
    text = text.replace(".", "").replace(",", ".")
    return float(text)


def read_rows(path: str | Path) -> tuple[list[PoRow], list[str]]:
    """Returns (rows, errors). Header row = first row containing the 3 columns."""
    ws = load_workbook(path, data_only=True).active
    cols: dict[str, int] = {}
    header_row = None
    for r in range(1, min(ws.max_row, 20) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            key = _match(str(ws.cell(r, c).value or ""))
            if key and key not in found:
                found[key] = c
        if len(found) == 3:
            cols, header_row = found, r
            break
    if not header_row:
        return [], ["No se encontraron las columnas NIT, VALOR y LIQUIDACION en las primeras 20 filas."]

    rows, errors = [], []
    for r in range(header_row + 1, ws.max_row + 1):
        nit = _clean_id(ws.cell(r, cols["nit"]).value)
        raw_amount = ws.cell(r, cols["amount"]).value
        liquidation = _clean_id(ws.cell(r, cols["liquidation"]).value)
        if not nit and raw_amount in (None, "") and not liquidation:
            continue
        try:
            amount = _to_amount(raw_amount)
            if amount <= 0:
                raise ValueError
        except (ValueError, TypeError):
            errors.append(f"Fila {r}: valor inválido '{raw_amount}'")
            continue
        if not nit or not liquidation:
            errors.append(f"Fila {r}: falta NIT o número de liquidación")
            continue
        rows.append(PoRow(r, nit, amount, liquidation))
    return rows, errors


def write_results(rows: list[PoRow], source: str | Path, out: Path | None = None) -> Path:
    source = Path(source)
    out = out or source.with_name(f"{source.stem}_resultado_{datetime.now():%Y%m%d_%H%M%S}.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultado"
    headers = ["Fila", "NIT", "Valor", "Liquidación", "Orden de compra", "Estado", "Mensaje"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    fills = {"OK": "C6EFCE", "ERROR": "FFC7CE", "REVISAR": "FFEB9C"}
    for x in rows:
        ws.append([x.row, x.nit, x.amount, x.liquidation, x.po_number, x.status, x.message])
        if x.status in fills:
            ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor=fills[x.status])
        ws.cell(ws.max_row, 3).number_format = "#,##0"
    for col, width in zip("ABCDEFG", (6, 14, 14, 14, 16, 10, 70)):
        ws.column_dimensions[col].width = width
    wb.save(out)
    return out


def write_template(path: str | Path) -> None:
    """Empty input template: same headers the reader expects, NIT/liquidation stored as text."""
    from openpyxl.comments import Comment
    wb = Workbook()
    ws = wb.active
    ws.title = "OCs"
    ws.append(["NIT PROVEEDOR", "VALOR OC", "NUMERO LIQUIDACION"])
    for cell, note in zip(ws[1], ("NIT o código del proveedor en SAP, sin puntos ni dígito de verificación.",
                                  "Valor total de la OC (sin IVA), solo números.",
                                  "Número de liquidación; se guarda como texto de la posición.")):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.comment = Comment(note, "Elogia")
    for col, width in zip("ABC", (20, 18, 24)):
        ws.column_dimensions[col].width = width
    for r in range(2, 502):
        ws.cell(r, 1).number_format = "@"
        ws.cell(r, 2).number_format = "#,##0"
        ws.cell(r, 3).number_format = "@"
    wb.save(path)
